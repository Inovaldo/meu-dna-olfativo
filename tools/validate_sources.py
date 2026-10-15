import json, pathlib, math
import openpyxl
from PIL import Image, ImageDraw
root=pathlib.Path(__file__).resolve().parents[1]
d=json.loads((root/'odores_157.json').read_text(encoding='utf-8'))
w=openpyxl.load_workbook(root/'mapa_olfativo_reformulado.xlsx',read_only=True,data_only=False)
summary={'jsonOdors':len(d['odors']),'families':d['families'],'sheets':[]}
for s in w:
 rows=list(s.values)
 summary['sheets'].append({'name':s.title,'rows':s.max_row,'columns':s.max_column,'sample':[[str(c) if c is not None else None for c in row] for row in rows[:7]]})
questionnaire=list(w.worksheets[1].values)
matched=[]
for o in d['odors']:
 found=[row for row in questionnaire if o['name'] in row]
 matched.append(bool(found))
summary['exactOdorNamesMatchedInWorkbook']=sum(matched)
summary['unmatchedNames']=[o['name'] for o,m in zip(d['odors'],matched) if not m]
(root/'artifacts/source-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
# Original geometric icon: three scent trails emerging from a seed.
im=Image.new('RGB',(1024,1024),'#235f53'); draw=ImageDraw.Draw(im)
draw.ellipse((425,680,599,854),fill='#f6edcf')
for shift in [-170,0,170]:
 points=[]
 for step in range(101):
  y=230+step*4.2
  x=512+shift+38*math.sin(step/100*math.pi*2)
  points.append((x,y))
 draw.line(points,fill='#f6edcf',width=38)
for folder,size in [('mdpi',48),('hdpi',72),('xhdpi',96),('xxhdpi',144),('xxxhdpi',192)]:
 im.resize((size,size),Image.Resampling.LANCZOS).save(root/f'android/app/src/main/res/mipmap-{folder}/ic_launcher.png')
ios=root/'ios/Runner/Assets.xcassets/AppIcon.appiconset'
for icon in json.loads((ios/'Contents.json').read_text())['images']:
 if 'filename' in icon:
  size=round(float(icon['size'].split('x')[0])*float(icon['scale'].rstrip('x')))
  im.resize((size,size),Image.Resampling.LANCZOS).save(ios/icon['filename'])
im.resize((192,192)).save(root/'web/icons/Icon-192.png');im.resize((512,512)).save(root/'web/icons/Icon-512.png')
im.resize((192,192)).save(root/'web/icons/Icon-maskable-192.png');im.resize((512,512)).save(root/'web/icons/Icon-maskable-512.png')
im.resize((32,32)).save(root/'web/favicon.png');im.save(root/'assets/icon-original.png')
print('Catalog verified:',len(d['odors']),len(d['families']),'exact workbook matches:',sum(matched))
